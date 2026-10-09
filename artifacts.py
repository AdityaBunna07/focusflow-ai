"""Local document-generation tools for FocusFlow execution actions."""

from __future__ import annotations

from datetime import datetime
from io import BytesIO
from typing import Any


def clean_title(plan: dict[str, Any]) -> str:
    return str(plan.get("summary") or "FocusFlow execution plan")[:110]


def grouped_schedule(plan: dict[str, Any]) -> dict[str, list[dict[str, Any]]]:
    groups: dict[str, list[dict[str, Any]]] = {}
    for item in plan.get("schedule", []):
        groups.setdefault(item["scheduled_date"], []).append(item)
    return groups


def create_pptx(plan: dict[str, Any]) -> bytes:
    """Generate a task-plan presentation locally; no external service is used."""
    from pptx import Presentation
    from pptx.util import Inches, Pt

    presentation = Presentation()
    title_slide = presentation.slides.add_slide(presentation.slide_layouts[0])
    title_slide.shapes.title.text = "FocusFlow AI"
    title_slide.placeholders[1].text = clean_title(plan)

    overview = presentation.slides.add_slide(presentation.slide_layouts[1])
    overview.shapes.title.text = "Execution overview"
    body = overview.placeholders[1].text_frame
    body.text = f"Daily capacity: {plan.get('hours_per_day', 0)} hours"
    for task in plan.get("tasks", []):
        paragraph = body.add_paragraph()
        deadline = f" — due {task['deadline']}" if task.get("deadline") else ""
        paragraph.text = f"{task['title']} ({task['estimated_minutes']} min){deadline}"
        paragraph.level = 0

    for day, items in grouped_schedule(plan).items():
        slide = presentation.slides.add_slide(presentation.slide_layouts[1])
        slide.shapes.title.text = f"Plan for {day}"
        body = slide.placeholders[1].text_frame
        body.clear()
        for item in items:
            paragraph = body.add_paragraph()
            paragraph.text = f"{item['title']} — {item['estimated_minutes']} minutes (priority {item.get('priority_score', '-')})"
            paragraph.font.size = Pt(20)

    output = BytesIO()
    presentation.save(output)
    return output.getvalue()


def create_pdf(plan: dict[str, Any]) -> bytes:
    """Generate a readable execution brief as a local PDF."""
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.units import cm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

    output = BytesIO()
    document = SimpleDocTemplate(output, pagesize=A4, rightMargin=1.8 * cm, leftMargin=1.8 * cm, topMargin=1.8 * cm, bottomMargin=1.8 * cm)
    styles = getSampleStyleSheet()
    story = [Paragraph("FocusFlow AI: Execution brief", styles["Title"]), Spacer(1, 12)]
    story.append(Paragraph(clean_title(plan), styles["BodyText"]))
    story.append(Spacer(1, 12))
    story.append(Paragraph(f"Generated: {datetime.now().strftime('%d %b %Y, %H:%M')}", styles["BodyText"]))
    story.append(Spacer(1, 16))
    for day, items in grouped_schedule(plan).items():
        story.append(Paragraph(day, styles["Heading2"]))
        for item in items:
            deadline = f"; due {item['deadline']}" if item.get("deadline") else ""
            text = f"<b>{item['title']}</b> — {item['estimated_minutes']} min; priority {item.get('priority_score', '-')}{deadline}"
            story.append(Paragraph(text, styles["BodyText"]))
            if item.get("notes"):
                story.append(Paragraph(item["notes"], styles["BodyText"]))
        story.append(Spacer(1, 10))
    document.build(story)
    return output.getvalue()


def create_docx(plan: dict[str, Any]) -> bytes:
    """Generate an editable Word execution report locally."""
    from docx import Document

    document = Document()
    document.add_heading("FocusFlow AI: Execution report", 0)
    document.add_paragraph(clean_title(plan))
    document.add_paragraph(f"Daily capacity: {plan.get('hours_per_day', 0)} hours")
    for day, items in grouped_schedule(plan).items():
        document.add_heading(day, level=1)
        for item in items:
            deadline = f"; due {item['deadline']}" if item.get("deadline") else ""
            document.add_paragraph(
                f"{item['title']} — {item['estimated_minutes']} minutes; priority {item.get('priority_score', '-')}{deadline}",
                style="List Bullet",
            )
    output = BytesIO()
    document.save(output)
    return output.getvalue()
